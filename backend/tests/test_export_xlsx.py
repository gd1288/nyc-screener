"""The workbook's formulas must reproduce the engine, not just report its answers.

Every assertion here evaluates the saved .xlsx with the `formulas` package and compares against
`engine.run()`. That is the only thing standing between "the export is a live model" and "the export
looks live and drifts the moment someone edits a cell".
"""

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.db import Base
from app.models import Neighborhood, NeighborhoodMetric, NeighborhoodScore
from app.services import MarketContext
from app.valuation import engine
from app.valuation import export_xlsx as export
from app.valuation.factors import PropertyProfile

formulas = pytest.importorskip("formulas")

SQUARE = {
    "type": "Polygon",
    "coordinates": [[[-74.02, 40.70], [-74.00, 40.70], [-74.00, 40.72], [-74.02, 40.72], [-74.02, 40.70]]],
}
DOLLAR = 1.0


@pytest.fixture(scope="module")
def ctx():
    engine_ = create_engine("sqlite://")
    Base.metadata.create_all(engine_)
    session = sessionmaker(bind=engine_)()
    session.add(
        Neighborhood(code="MN0101", name="Test", borough="Manhattan", residential=True, geometry=SQUARE, area_km2=4)
    )
    session.add(NeighborhoodScore(nta_code="MN0101", score=65, rank=1, pillars={}, coverage=1.0))
    session.add(NeighborhoodMetric(nta_code="MN0101", metric="zhvi_cagr_10y", value=0.05, source="z", as_of="2025"))
    session.add(NeighborhoodMetric(nta_code="MN0101", metric="zori_rent", value=3500, source="z", as_of="2025"))
    session.commit()
    return MarketContext.load(session)


PROFILE = PropertyProfile(price=1_200_000, nta_code="MN0101", sqft=800, bedrooms=2, rent_estimate=5_200)


def evaluate(tmp_path, ctx, overrides=None, name="book.xlsx"):
    """Save the workbook and return {(SHEET, CELL): value} as Excel itself would recalculate it."""
    path = tmp_path / name
    export.build_workbook(PROFILE, ctx, overrides or {}, label="Test", monte_carlo_runs=100, seed=1).save(path)
    solution = formulas.ExcelModel().loads(str(path)).finish().calculate()
    values = {}
    for key, value in solution.items():
        if "'!" not in key:
            continue
        sheet, cell = key.rsplit("'!", 1)
        try:
            values[(sheet.split("]")[-1], cell)] = value.value[0, 0]
        except (AttributeError, IndexError, TypeError):
            continue
    return values


def cash_flow(values, row, column):
    return values[("CASH FLOWS", f"{column}{row}")]


def _threadsafe_memory_engine():
    """TestClient serves requests on its own thread, and a plain in-memory SQLite connection is
    bound to the thread that opened it — so the request handler would see a different, empty DB."""
    from sqlalchemy.pool import StaticPool

    return create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)


@pytest.mark.parametrize(
    "overrides, case",
    [
        ({}, "defaults"),
        # A loan shorter than the hold period is the case that actually exercises the
        # debt-service-stops branch in `_project`. At the defaults (30y loan, 20y hold) a naive
        # `=monthly_payment*12` formula matches perfectly while being wrong, so the defaults alone
        # cannot tell a correct sheet from an incorrect one.
        ({"loan_years": 15}, "loan shorter than the 20-year hold"),
    ],
)
def test_cash_flow_rows_match_the_engine_to_the_dollar(tmp_path, ctx, overrides, case):
    values = evaluate(tmp_path, ctx, overrides, name=f"{case.replace(' ', '_')}.xlsx")
    expected = engine.run(PROFILE, ctx, overrides)["projections"]["base"]["20"]["path"]

    for year_data in expected:
        row = year_data["year"] + 2
        assert cash_flow(values, row, "F") == pytest.approx(year_data["noi"], abs=DOLLAR), (
            f"NOI year {year_data['year']}"
        )
        assert -cash_flow(values, row, "G") == pytest.approx(year_data["mortgage"], abs=DOLLAR), (
            f"debt service year {year_data['year']} ({case})"
        )
        assert cash_flow(values, row, "H") == pytest.approx(year_data["cash_flow"], abs=DOLLAR)
        assert cash_flow(values, row, "I") == pytest.approx(year_data["cumulative_cash_flow"], abs=DOLLAR)
        assert cash_flow(values, row, "K") == pytest.approx(year_data["property_value"], abs=DOLLAR)
        assert cash_flow(values, row, "L") == pytest.approx(year_data["loan_balance"], abs=DOLLAR)
        assert cash_flow(values, row, "M") == pytest.approx(year_data["equity"], abs=DOLLAR)


def test_debt_service_actually_stops_when_the_loan_is_repaid(tmp_path, ctx):
    """Guards the guard: if this ever stopped being true, the parametrized case above would be
    testing the same code path twice and quietly proving nothing."""
    values = evaluate(tmp_path, ctx, {"loan_years": 15}, name="payoff.xlsx")
    assert cash_flow(values, 15 + 2, "G") < 0  # year 15: still paying
    assert cash_flow(values, 16 + 2, "G") == pytest.approx(0, abs=1e-9)  # year 16: loan is gone


def test_irr_and_equity_multiple_match_the_engine(tmp_path, ctx):
    values = evaluate(tmp_path, ctx)
    projections = engine.run(PROFILE, ctx, {})["projections"]["base"]
    results_row = export.MAX_YEARS + 4
    for i, horizon in enumerate(export.HORIZONS):
        assert values[("CASH FLOWS", f"B{results_row + 1 + i}")] == pytest.approx(
            projections[str(horizon)]["irr"], abs=1e-4
        )
        assert values[("CASH FLOWS", f"D{results_row + 1 + i}")] == pytest.approx(
            projections[str(horizon)]["equity_multiple"], abs=0.01
        )


def test_cash_invested_and_purchase_costs_match_the_engine(tmp_path, ctx):
    values = evaluate(tmp_path, ctx)
    analysis = engine.run(PROFILE, ctx, {})
    assert values[("CASH FLOWS", "N2")] == pytest.approx(-analysis["cash_invested"], abs=DOLLAR)
    assert values[("CASH FLOWS", "O2")] == pytest.approx(-analysis["cash_invested"], abs=DOLLAR)


def test_editing_an_assumption_recalculates_the_projection(tmp_path, ctx):
    """The point of the export. A higher mortgage rate has to move cash flow in the workbook itself,
    not just in whatever the engine happened to write at export time."""
    base = evaluate(tmp_path, ctx, {"interest_rate": 0.05}, name="low.xlsx")
    high = evaluate(tmp_path, ctx, {"interest_rate": 0.09}, name="high.xlsx")
    assert cash_flow(high, 3, "G") < cash_flow(base, 3, "G")  # more negative: a bigger payment
    assert cash_flow(high, 3, "H") < cash_flow(base, 3, "H")


def test_mansion_tax_re_rates_from_the_price_cell(tmp_path, ctx):
    """The bracket schedule is a live nested IF, so it is generated from `inv.MANSION_TAX_BRACKETS`
    and must agree with the Python lookup at every bracket edge."""
    from app.scoring import investment as inv

    formula = export._mansion_rate_formula("price")
    for price in (900_000, 1_000_000, 2_500_000, 3_000_000, 5_000_000, 12_000_000, 26_000_000):
        evaluated = formulas.Parser().ast(f"={formula.replace('price', str(price))}")[1].compile()()
        assert float(evaluated) == pytest.approx(inv.mansion_tax_rate(price)), f"bracket at {price}"


def test_workbook_has_the_expected_sheets_and_named_cells(tmp_path, ctx):
    import openpyxl

    path = tmp_path / "sheets.xlsx"
    export.build_workbook(PROFILE, ctx, {}, label="Test", monte_carlo_runs=100, seed=1).save(path)
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["Assumptions", "Cash Flows", "Scenarios", "Sensitivity", "Monte Carlo"]
    for name in ("price", "interest_rate", "cash_invested", "opex_monthly", "irr_10"):
        assert name in wb.defined_names
    assert wb["Cash Flows"].freeze_panes == "B2"


def test_cash_flow_cells_are_formulas_not_baked_numbers(tmp_path, ctx):
    import openpyxl

    path = tmp_path / "live.xlsx"
    export.build_workbook(PROFILE, ctx, {}, label="Test", monte_carlo_runs=100, seed=1).save(path)
    ws = openpyxl.load_workbook(path)["Cash Flows"]
    for row in range(3, export.MAX_YEARS + 3):
        for column in ("B", "F", "H", "K", "L"):
            assert str(ws[f"{column}{row}"].value).startswith("="), f"{column}{row} is not a formula"


def test_export_endpoint_returns_a_real_workbook_with_a_download_filename(tmp_path):
    """The route is the deliverable users actually hit, so cover it separately from the builder."""
    import io

    import openpyxl
    from fastapi.testclient import TestClient
    from sqlalchemy.orm import sessionmaker

    from app.api import valuation as valuation_api
    from app.main import app
    from app.models import ValuationProperty

    db_engine = _threadsafe_memory_engine()
    Base.metadata.create_all(db_engine)
    Session = sessionmaker(bind=db_engine)
    with Session() as session:
        session.add(
            Neighborhood(code="MN0101", name="Test", borough="Manhattan", residential=True, geometry=SQUARE, area_km2=4)
        )
        session.add(
            ValuationProperty(
                label='My "Flat" / 3A', property_type="condo", nta_code="MN0101", price=1_200_000, rent_estimate=5_200
            )
        )
        session.commit()

    app.dependency_overrides[valuation_api.db] = lambda: Session()
    try:
        response = TestClient(app).get("/api/valuation/properties/1/export.xlsx")
    finally:
        app.dependency_overrides.clear()

    assert response.status_code == 200
    assert response.headers["content-type"] == valuation_api.XLSX_MEDIA_TYPE
    # Quotes and slashes from a user-typed label must not reach the header raw.
    disposition = response.headers["content-disposition"]
    assert disposition == 'attachment; filename="My-_Flat_-_-3A.xlsx"'
    workbook = openpyxl.load_workbook(io.BytesIO(response.content))
    assert "Cash Flows" in workbook.sheetnames


def test_export_endpoint_404s_for_an_unknown_property():
    from fastapi.testclient import TestClient
    from sqlalchemy.orm import sessionmaker

    from app.api import valuation as valuation_api
    from app.main import app

    db_engine = _threadsafe_memory_engine()
    Base.metadata.create_all(db_engine)
    app.dependency_overrides[valuation_api.db] = lambda: sessionmaker(bind=db_engine)()
    try:
        assert TestClient(app).get("/api/valuation/properties/999/export.xlsx").status_code == 404
    finally:
        app.dependency_overrides.clear()
