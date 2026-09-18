"""Excel export of a valuation run — a live model, not a picture of one.

The Cash Flows sheet is built from formulas that reference named cells on the Assumptions sheet, so
changing a rate or a growth rate in Excel recalculates the whole projection, IRR included. Dumping
the engine's computed numbers would have been far less code and would have produced a workbook that
silently lies the moment anyone edits an input.

Keeping two implementations of the same cash-flow math (Python and Excel) honest is the job of
`tests/test_export_xlsx.py`, which evaluates the workbook with the `formulas` package and asserts
the results match `engine.run()` to the dollar — including a case with a loan shorter than the hold
period, because the debt-service-stops branch never fires at the defaults.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from app.scoring import investment as inv
from app.services import MarketContext
from app.valuation import engine
from app.valuation import scenarios as scenarios_mod
from app.valuation.factors import PropertyProfile

MAX_YEARS = 20
HORIZONS = (10, 20)

MONEY = "#,##0"
PERCENT = "0.00%"
RATIO = "0.00"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
SECTION_FONT = Font(bold=True, color="1F3864")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")  # yellow = safe to edit, the usual model convention


def _mansion_rate_formula(price_ref: str = "price") -> str:
    """The NYS mansion-tax schedule as a live nested IF, so the tax re-rates when price is edited.

    Generated from `inv.MANSION_TAX_BRACKETS` rather than transcribed — a hand-copied bracket table
    is a second source of truth that would drift the first time the schedule changes.
    """
    expression = "0"
    for threshold, rate in reversed(inv.MANSION_TAX_BRACKETS):
        expression = f"IF({price_ref}>={threshold},{rate},{expression})"
    return expression


class _Assumptions:
    """Writes the Assumptions sheet and registers a defined name for every cell, so formulas on the
    other sheets read as `price*(1-down_payment_pct)` rather than `$B$7*(1-$B$12)`."""

    def __init__(self, workbook: Workbook):
        self.wb = workbook
        self.ws = workbook.create_sheet("Assumptions")
        self.ws.column_dimensions["A"].width = 34
        self.ws.column_dimensions["B"].width = 18
        self.ws.column_dimensions["C"].width = 52
        self.row = 1

    def title(self, text: str) -> None:
        cell = self.ws.cell(self.row, 1, text)
        cell.font = Font(bold=True, size=14)
        self.row += 2

    def section(self, text: str) -> None:
        cell = self.ws.cell(self.row, 1, text)
        cell.font = SECTION_FONT
        self.row += 1

    def value(self, name: str, label: str, value, fmt: str, note: str = "", editable: bool = True) -> None:
        self.ws.cell(self.row, 1, label)
        cell = self.ws.cell(self.row, 2, value)
        cell.number_format = fmt
        if editable:
            cell.fill = INPUT_FILL
        if note:
            self.ws.cell(self.row, 3, note).font = Font(italic=True, size=9, color="808080")
        self.wb.defined_names.add(DefinedName(name, attr_text=f"Assumptions!$B${self.row}"))
        self.row += 1

    def formula(self, name: str, label: str, formula: str, fmt: str, note: str = "") -> None:
        self.value(name, label, formula, fmt, note, editable=False)

    def blank(self) -> None:
        self.row += 1


def _write_assumptions(wb: Workbook, profile: PropertyProfile, analysis: dict, label: str) -> None:
    a = analysis["assumptions"]
    exact = analysis["exact"]
    sheet = _Assumptions(wb)
    sheet.title(f"Valuation — {label}")

    sheet.section("Property")
    sheet.value("price", "Purchase price", profile.price, MONEY)
    sheet.value("rent_monthly", "Monthly rent", exact["rent"], MONEY, f"basis: {analysis['rent_basis']}")
    sheet.value("common_charges_monthly", "Monthly common charges", exact["common_charges"], MONEY)
    sheet.value("property_taxes_monthly", "Monthly property taxes", exact["property_taxes"], MONEY)
    sheet.blank()

    sheet.section("Financing")
    sheet.value("down_payment_pct", "Down payment %", a["down_payment_pct"], PERCENT)
    sheet.value("interest_rate", "Mortgage rate", a["interest_rate"], PERCENT)
    sheet.value("loan_years", "Loan term (years)", a["loan_years"], "0")
    sheet.formula("loan_amount", "Loan amount", "=price*(1-down_payment_pct)", MONEY)
    sheet.formula(
        "monthly_payment",
        "Monthly payment",
        "=IF(loan_amount<=0,0,PMT(interest_rate/12,loan_years*12,-loan_amount))",
        MONEY,
    )
    sheet.blank()

    sheet.section("Operating")
    sheet.value("vacancy_pct", "Vacancy", a["vacancy_pct"], PERCENT)
    sheet.value("management_pct", "Management fee (% of rent)", a["management_pct"], PERCENT)
    sheet.value("maintenance_monthly", "Maintenance (monthly)", a["maintenance_monthly"], MONEY)
    sheet.value("insurance_monthly", "Insurance (monthly)", a["insurance_monthly"], MONEY)
    sheet.formula(
        "opex_monthly",
        "Total monthly operating expense",
        "=common_charges_monthly+property_taxes_monthly+insurance_monthly+maintenance_monthly"
        "+rent_monthly*management_pct",
        MONEY,
        "management is charged on year-1 rent, then grown with expenses",
    )
    sheet.blank()

    sheet.section("Growth and exit")
    sheet.value("rent_growth", "Rent growth", a["rent_growth"], PERCENT)
    sheet.value("expense_growth", "Expense growth", a["expense_growth"], PERCENT)
    sheet.value("appreciation", "Appreciation", analysis["appreciation"]["base"], PERCENT, "base case")
    sheet.value("exit_cost_pct", "Selling costs", a["exit_cost_pct"], PERCENT)
    sheet.blank()

    sheet.section("Purchase costs")
    sheet.formula("mansion_tax", "Mansion tax", f"=price*({_mansion_rate_formula()})", MONEY, "NYS schedule")
    sheet.formula(
        "mortgage_recording_tax",
        "Mortgage recording tax",
        "=IF(loan_amount>0,loan_amount*IF(loan_amount>=500000,0.01925,0.018),0)",
        MONEY,
    )
    sheet.formula("title_insurance", "Title insurance", "=price*0.0045", MONEY)
    sheet.formula("attorney_and_bank_fees", "Attorney and bank fees", "=IF(loan_amount>0,5000,3000)", MONEY)
    components = ["mansion_tax", "mortgage_recording_tax", "title_insurance", "attorney_and_bank_fees"]
    if a["new_development"]:
        sheet.formula(
            "nyc_transfer_tax", "NYC transfer tax", "=price*IF(price>500000,0.01425,0.01)", MONEY, "sponsor sale"
        )
        sheet.formula(
            "nys_transfer_tax", "NYS transfer tax", "=price*IF(price>=3000000,0.0065,0.004)", MONEY, "sponsor sale"
        )
        components += ["nyc_transfer_tax", "nys_transfer_tax"]
    sheet.formula("purchase_costs_total", "Total purchase costs", f"={'+'.join(components)}", MONEY)
    sheet.formula("cash_invested", "Cash invested", "=price*down_payment_pct+purchase_costs_total", MONEY)

    sheet.ws.freeze_panes = "A3"


CASH_FLOW_COLUMNS = [
    ("Year", 8),
    ("Gross rent", 14),
    ("Vacancy loss", 14),
    ("Effective rent", 15),
    ("Operating expenses", 18),
    ("NOI", 14),
    ("Debt service", 14),
    ("Operating cash flow", 18),
    ("Cumulative cash flow", 19),
    ("DSCR", 9),
    ("Property value", 16),
    ("Loan balance", 15),
    ("Equity", 15),
    ("Net CF (10y exit)", 17),
    ("Net CF (20y exit)", 17),
]


def _write_cash_flows(wb: Workbook) -> None:
    ws = wb.create_sheet("Cash Flows")
    for i, (heading, width) in enumerate(CASH_FLOW_COLUMNS, start=1):
        cell = ws.cell(1, i, heading)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width

    # Year 0 carries only the initial outflow; the operating columns start at year 1.
    ws.cell(2, 1, 0)
    for col in (14, 15):
        ws.cell(2, col, "=-cash_invested").number_format = MONEY

    for year in range(1, MAX_YEARS + 1):
        row = year + 2
        prior = row - 1
        grown = f"(1+rent_growth)^{year - 1}"
        ws.cell(row, 1, year)
        ws.cell(row, 2, f"=rent_monthly*12*{grown}")
        ws.cell(row, 3, f"=-B{row}*vacancy_pct")
        ws.cell(row, 4, f"=B{row}+C{row}")
        ws.cell(row, 5, f"=-opex_monthly*12*(1+expense_growth)^{year - 1}")
        ws.cell(row, 6, f"=D{row}+E{row}")
        # Debt service stops once the loan is repaid, so a hold longer than the loan term isn't
        # charged a payment against a zero balance.
        ws.cell(row, 7, f"=-monthly_payment*MIN(12,MAX(loan_years*12-{(year - 1) * 12},0))")
        ws.cell(row, 8, f"=F{row}+G{row}")
        ws.cell(row, 9, f"=H{row}+I{prior}" if year > 1 else f"=H{row}")
        ws.cell(row, 10, f'=IF(G{row}=0,"",-F{row}/G{row})')
        ws.cell(row, 11, f"=price*(1+appreciation)^{year}")
        ws.cell(row, 12, f"=MAX(FV(interest_rate/12,{year * 12},monthly_payment,-loan_amount),0)")
        ws.cell(row, 13, f"=K{row}-L{row}")
        for horizon, col in zip(HORIZONS, (14, 15), strict=True):
            if year < horizon:
                ws.cell(row, col, f"=H{row}")
            elif year == horizon:
                ws.cell(row, col, f"=H{row}+K{row}*(1-exit_cost_pct)-L{row}")

        for col in (2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 13, 14, 15):
            ws.cell(row, col).number_format = MONEY
        ws.cell(row, 10).number_format = RATIO

    results_row = MAX_YEARS + 4
    ws.cell(results_row, 1, "Results").font = SECTION_FONT
    for i, (horizon, col) in enumerate(zip(HORIZONS, (14, 15), strict=True)):
        letter = get_column_letter(col)
        last = horizon + 2
        ws.cell(results_row + 1 + i, 1, f"IRR ({horizon}y)").font = Font(bold=True)
        irr_cell = ws.cell(results_row + 1 + i, 2, f"=IRR({letter}2:{letter}{last})")
        irr_cell.number_format = PERCENT
        wb.defined_names.add(DefinedName(f"irr_{horizon}", attr_text=f"'Cash Flows'!$B${results_row + 1 + i}"))
        ws.cell(results_row + 1 + i, 3, f"Equity multiple ({horizon}y)")
        multiple = ws.cell(results_row + 1 + i, 4, f"=SUM({letter}3:{letter}{last})/cash_invested")
        multiple.number_format = RATIO

    ws.freeze_panes = "B2"


def _write_scenarios(wb: Workbook, results: list[scenarios_mod.ScenarioResult]) -> None:
    ws = wb.create_sheet("Scenarios")
    headings = [
        "Scenario",
        "Mortgage rate",
        "Appreciation",
        "Rent growth",
        "Vacancy",
        "Cap rate",
        "Cash on cash",
        "Monthly cash flow",
        "IRR (10y)",
        "IRR (20y)",
        "Equity multiple (10y)",
        "Payback year",
    ]
    for i, heading in enumerate(headings, start=1):
        cell = ws.cell(1, i, heading)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 17

    for row, result in enumerate(results, start=2):
        a = result.assumptions
        ws.cell(row, 1, result.name)
        for col, (value, fmt) in enumerate(
            [
                (a["interest_rate"], PERCENT),
                (a["appreciation_override"], PERCENT),
                (a["rent_growth"], PERCENT),
                (a["vacancy_pct"], PERCENT),
                (result.cap_rate, PERCENT),
                (result.cash_on_cash, PERCENT),
                (result.monthly_cash_flow, MONEY),
                (result.irr_10, PERCENT),
                (result.irr_20, PERCENT),
                (result.equity_multiple_10, RATIO),
                (result.payback_year_10, "0"),
            ],
            start=2,
        ):
            ws.cell(row, col, value).number_format = fmt
    ws.freeze_panes = "B2"


def _write_sensitivity(wb: Workbook, rows: list[dict]) -> None:
    ws = wb.create_sheet("Sensitivity")
    headings = ["Factor", "Low value", "High value", "IRR at low", "IRR at high", "Swing", "Backed by a source"]
    for i, heading in enumerate(headings, start=1):
        cell = ws.cell(1, i, heading)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 18

    for row, item in enumerate(rows, start=2):
        ws.cell(row, 1, item["label"])
        ws.cell(row, 2, item["low_value"]).number_format = PERCENT
        ws.cell(row, 3, item["high_value"]).number_format = PERCENT
        ws.cell(row, 4, item["low_irr"]).number_format = PERCENT
        ws.cell(row, 5, item["high_irr"]).number_format = PERCENT
        ws.cell(row, 6, item["swing"]).number_format = PERCENT
        ws.cell(row, 7, "yes" if item["backed_by_source"] else "no — engine default")
    ws.freeze_panes = "A2"


def _write_monte_carlo(wb: Workbook, mc: dict) -> None:
    ws = wb.create_sheet("Monte Carlo")
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 16
    ws.cell(1, 1, "Monte Carlo").font = Font(bold=True, size=12)
    summary = [
        ("Runs", mc["n"], "0"),
        ("Runs with a solvable IRR", mc["valid_n"], "0"),
        ("P10 IRR", mc["p10"], PERCENT),
        ("P50 IRR", mc["p50"], PERCENT),
        ("P90 IRR", mc["p90"], PERCENT),
        ("Probability of a negative IRR", mc["prob_loss"], PERCENT),
    ]
    for i, (label, value, fmt) in enumerate(summary, start=2):
        ws.cell(i, 1, label)
        ws.cell(i, 2, value).number_format = fmt

    # Only factors with a real P10/P90 are sampled; the rest are held at their default, and saying
    # so here stops the distribution being read as wider evidence than it is.
    ws.cell(9, 1, "Randomized factors").font = SECTION_FONT
    ws.cell(9, 2, ", ".join(mc["randomized_factors"]) or "none")
    ws.cell(10, 1, "Held at default").font = SECTION_FONT
    ws.cell(10, 2, ", ".join(mc["held_at_default"]) or "none")

    ws.cell(12, 1, "Histogram").font = SECTION_FONT
    for i, heading in enumerate(["IRR from", "IRR to", "Count"], start=1):
        cell = ws.cell(13, i, heading)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for row, item in enumerate(mc["histogram"], start=14):
        ws.cell(row, 1, item["bin_start"]).number_format = PERCENT
        ws.cell(row, 2, item["bin_end"]).number_format = PERCENT
        ws.cell(row, 3, item["count"])


def build_workbook(
    profile: PropertyProfile,
    ctx: MarketContext,
    overrides: dict[str, float] | None = None,
    label: str = "Property",
    monte_carlo_runs: int = 500,
    seed: int | None = 0,
) -> Workbook:
    overrides = overrides or {}
    analysis = engine.run(profile, ctx, overrides)
    wb = Workbook()
    wb.remove(wb.active)
    _write_assumptions(wb, profile, analysis, label)
    _write_cash_flows(wb)
    _write_scenarios(wb, scenarios_mod.compare(profile, ctx, overrides))
    _write_sensitivity(wb, scenarios_mod.sensitivity(profile, ctx, overrides))
    _write_monte_carlo(wb, scenarios_mod.monte_carlo(profile, ctx, overrides, n=monte_carlo_runs, seed=seed))
    return wb


def to_bytes(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
